#!/usr/bin/env python3
"""
Update diagrams in Excel workbook after SVG fixes.
Re-inserts the fixed SVG diagrams (converted to PNG) into the existing Excel workbook.
Preserves all formulas, formatting, and layout - only replaces the images.
"""

import os
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image
import io
import cairosvg

# Set up paths
repo_root = Path(__file__).parent.parent
docs_img = repo_root / "docs" / "img"
templates_dir = repo_root / "templates"
workbook_path = templates_dir / "worksheet.xlsx"

# Define which diagrams go on which sheets and at what size
diagram_mappings = {
    "Foundation": {
        "svg": docs_img / "foundation-plan-section.svg",
        "anchor": "D2",
        "width": 700
    },
    "Roof Takeoff": {
        "svg": docs_img / "roof-plan-and-section.svg",
        "anchor": "P2",
        "width": 700
    }
}

def svg_to_png(svg_path, width_px):
    """Convert SVG to PNG with specified width, maintaining aspect ratio."""
    print(f"Converting {svg_path.name} to PNG at {width_px}px width...")
    
    # Read SVG file
    with open(svg_path, 'rb') as f:
        svg_data = f.read()
    
    # Convert SVG to PNG using cairosvg
    png_data = cairosvg.svg2png(bytestring=svg_data, output_width=width_px)
    
    # Load with PIL to get image object
    img = Image.open(io.BytesIO(png_data))
    
    return img

def update_excel_diagrams():
    """Update diagrams in the Excel workbook."""
    
    if not workbook_path.exists():
        print(f"Error: Workbook not found at {workbook_path}")
        return False
    
    print(f"\nLoading workbook: {workbook_path}")
    wb = load_workbook(workbook_path)
    
    for sheet_name, config in diagram_mappings.items():
        if sheet_name not in wb.sheetnames:
            print(f"Warning: Sheet '{sheet_name}' not found, skipping...")
            continue
        
        ws = wb[sheet_name]
        svg_path = config["svg"]
        
        if not svg_path.exists():
            print(f"Warning: SVG not found at {svg_path}, skipping...")
            continue
        
        print(f"\nProcessing sheet: {sheet_name}")
        
        # Remove existing images from the sheet at this anchor
        # We need to remove and re-add to update the image
        images_to_remove = []
        for img in ws._images:
            if hasattr(img, 'anchor') and hasattr(img.anchor, '_from'):
                anchor_cell = img.anchor._from
                if hasattr(anchor_cell, 'col'):
                    # Check if this image is at our target anchor
                    from openpyxl.utils import get_column_letter
                    img_col = get_column_letter(anchor_cell.col + 1)
                    img_row = anchor_cell.row + 1
                    target_anchor = config["anchor"]
                    
                    # Extract column letter and row from target
                    import re
                    match = re.match(r'([A-Z]+)(\d+)', target_anchor)
                    if match:
                        target_col, target_row = match.groups()
                        if img_col == target_col and img_row == int(target_row):
                            images_to_remove.append(img)
                            print(f"  Found existing image at {target_anchor}, will replace")
        
        # Remove old images
        for img in images_to_remove:
            ws._images.remove(img)
        
        # Convert SVG to PNG
        pil_img = svg_to_png(svg_path, config["width"])
        
        # Save to BytesIO
        img_byte_arr = io.BytesIO()
        pil_img.save(img_byte_arr, format='PNG')
        img_byte_arr.seek(0)
        
        # Create openpyxl Image
        xl_img = XLImage(img_byte_arr)
        xl_img.anchor = config["anchor"]
        
        # Add to worksheet
        ws.add_image(xl_img)
        print(f"  ✓ Inserted updated {svg_path.name} at {config['anchor']} ({pil_img.width}x{pil_img.height}px)")
    
    # Save workbook
    print(f"\nSaving workbook...")
    wb.save(workbook_path)
    print(f"✓ Workbook saved successfully: {workbook_path}")
    
    # Verify no merged cells
    merged_count = 0
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        merged_count += len(ws.merged_cells.ranges)
    
    print(f"\n✓ Validation: {merged_count} merged cells (should be 0)")
    
    if merged_count > 0:
        print("WARNING: Workbook contains merged cells!")
        return False
    
    return True

if __name__ == "__main__":
    print("=" * 60)
    print("UPDATING DIAGRAMS IN EXCEL WORKBOOK")
    print("=" * 60)
    
    success = update_excel_diagrams()
    
    if success:
        print("\n" + "=" * 60)
        print("SUCCESS: Diagrams updated in Excel workbook")
        print("=" * 60)
        sys.exit(0)
    else:
        print("\n" + "=" * 60)
        print("FAILED: Error updating diagrams")
        print("=" * 60)
        sys.exit(1)
